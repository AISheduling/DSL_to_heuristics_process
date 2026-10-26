"""
Индикаторы проекта.

1. Сети, посчитанные вручную: цепочка, параллельная сеть, ромб, избыточная длинная дуга.
2. Инварианты: перенумерация, фиктивные работы, диапазон [0, 1].
3. Диапазоны min/max по J30/J60/J120 из guo2021prselection, табл. 3–4 (если PSPLIB скачан).
4. Построчная сверка 11 индикаторов с RCPLIB на J30/J60/J120 (если скачаны PSPLIB и RCPLIB,
   data/psplib/SOURCES.md, data/rcplib/SOURCES.md).
"""
from __future__ import annotations

from pathlib import Path

import pytest

from indicators import (IndicatorError, compute_indicators, flat_indicators,
                        network_from_instance)

ROOT = Path(__file__).resolve().parents[1]
PSPLIB_DIR = ROOT / "data" / "psplib" / "instances"


def make_instance(durations: dict[int, int], arcs: list[tuple[int, int]],
                  demand: dict[int, dict[int, int]] | None = None,
                  capacity: dict[int, int] | None = None, dummies: bool = True) -> dict:
    """Инстанс в формате data/processed; при dummies=True добавляются source 0 и sink 99."""
    demand = demand or {}
    capacity = capacity or {1: 10}
    ids = sorted(durations)
    succ = {i: [j for a, j in arcs if a == i] for i in ids}
    jobs = []
    if dummies:
        roots = [i for i in ids if not any(j == i for _, j in arcs)]
        leaves = [i for i in ids if not succ[i]]
        jobs.append({"id": 0, "duration": 0, "successors": roots, "resource_requirements": []})
        for i in leaves:
            succ[i] = [99]
    for i in ids:
        jobs.append({"id": i, "duration": durations[i], "successors": succ[i],
                     "resource_requirements": [{"resource_id": k, "amount": a}
                                               for k, a in demand.get(i, {}).items()]})
    if dummies:
        jobs.append({"id": 99, "duration": 0, "successors": [], "resource_requirements": []})
    return {"resources": [{"id": k, "name": f"R{k}", "capacity": c, "type": "renewable"}
                          for k, c in capacity.items()],
            "jobs": jobs}


# Примеры, посчитанные вручную

def test_chain():
    ind = flat_indicators(make_instance({1: 2, 2: 3, 3: 1, 4: 4}, [(1, 2), (2, 3), (3, 4)]))
    assert ind["n"] == 4
    assert ind["cnc"] == pytest.approx(0.75)
    assert ind["os"] == 1
    assert ind["sp"] == 1
    assert ind["ad"] == 0
    assert ind["la"] == 1
    assert ind["i5"] == 1
    assert ind["tf"] == 0


def test_parallel():
    ind = flat_indicators(make_instance({1: 2, 2: 3, 3: 1, 4: 4}, []))
    assert ind["cnc"] == 0
    assert ind["os"] == 0
    assert ind["sp"] == 0
    assert ind["ad"] == 0
    assert ind["tf"] == 0


def test_diamond():
    # a, b -> c: уровни [2, 1], w̄ = 1.5, alpha_w = 1, alpha_max = 1
    ind = flat_indicators(make_instance({1: 1, 2: 1, 3: 1}, [(1, 3), (2, 3)]))
    assert ind["cnc"] == pytest.approx(2 / 3)
    assert ind["os"] == pytest.approx(2 / 3)
    assert ind["sp"] == pytest.approx(0.5)
    assert ind["ad"] == pytest.approx(1)
    assert ind["la"] == pytest.approx(1)
    assert ind["i5"] == pytest.approx(1)
    assert ind["tf"] == 0


def test_redundant_long_arc():
    # 1->2->3 плюс 1->3: дуга длины 2 при m = 3 весит (m − l − 1)/(m − 2) = 0
    without = flat_indicators(make_instance({1: 1, 2: 1, 3: 1}, [(1, 2), (2, 3)]))
    ind = flat_indicators(make_instance({1: 1, 2: 1, 3: 1}, [(1, 2), (2, 3), (1, 3)]))
    assert ind["os"] == without["os"] == 1
    assert ind["la"] == 1
    assert without["i5"] == 1
    assert ind["i5"] == 0


def test_topological_float():
    # 1->2->3 и отдельная 4: m = 3, у 4 запас RL − PL = 3 − 1 = 2; TF = 2 / ((m−1)(n−m)) = 1
    ind = flat_indicators(make_instance({1: 1, 2: 1, 3: 1, 4: 1}, [(1, 2), (2, 3)]))
    assert ind["sp"] == pytest.approx(2 / 3)
    assert ind["tf"] == pytest.approx(1)


def test_resource_strength_bounds():
    # Две параллельные работы по 2 единицы: r_min = 2, пик раннего расписания r_max = 4
    def rs(cap):
        inst = make_instance({1: 1, 2: 1}, [], demand={1: {1: 2}, 2: {1: 2}}, capacity={1: cap})
        return flat_indicators(inst)["rs"]
    assert rs(4) == 1
    assert rs(3) == pytest.approx(0.5)
    assert rs(2) == 0


def test_resource_factor_use_constrainedness():
    inst = make_instance({1: 1, 2: 1}, [(1, 2)], demand={1: {1: 2}, 2: {1: 4, 2: 1}},
                         capacity={1: 8, 2: 2})
    ind = flat_indicators(inst)
    assert ind["rf"] == pytest.approx(3 / 4)
    assert ind["ru"] == pytest.approx(1.5)
    # RC проекта — среднее по парам (работа, ресурс): 2/8, 4/8, 1/2; по ресурсам: RC_1 = 3/8, RC_2 = 1/2
    assert ind["rc"] == pytest.approx((2 / 8 + 4 / 8 + 1 / 2) / 3)
    per_resource = compute_indicators(inst)["resource"]["per_resource"]
    assert per_resource[1]["rc"] == pytest.approx(3 / 8)
    assert per_resource[2]["rc"] == pytest.approx(1 / 2)


def test_unused_resource_excluded_from_rs_rc():
    inst = make_instance({1: 1}, [], demand={1: {1: 2}}, capacity={1: 4, 2: 5})
    res = compute_indicators(inst)["resource"]
    assert res["per_resource"][2] == {"rs": None, "rc": None}
    assert res["rc"] == pytest.approx(0.5)
    assert res["rf"] == pytest.approx(0.5)


def test_extended_giant_job_and_bottleneck():
    # 1 — короткая, от нее зависят все; 2 — длиннее всех остальных вместе
    inst = make_instance({1: 1, 2: 20, 3: 2, 4: 2, 5: 2}, [(1, 2), (1, 3), (1, 4), (1, 5)])
    ext = compute_indicators(inst)["extended"]
    assert ext["max_duration_share"] == pytest.approx(20 / 27)
    assert ext["max_descendant_share"] == 1
    assert ext["cpm_length"] == 21


# Инварианты и область применимости

def test_dummies_do_not_change_values():
    args = ({1: 2, 2: 3, 3: 1, 4: 4}, [(1, 2), (1, 3), (2, 4)])
    kw = {"demand": {1: {1: 3}, 3: {1: 5}}}
    assert (flat_indicators(make_instance(*args, **kw, dummies=True))
            == flat_indicators(make_instance(*args, **kw, dummies=False)))


def test_renumbering_does_not_change_values():
    durations = {1: 2, 2: 3, 3: 1, 4: 4, 5: 2}
    arcs = [(1, 2), (1, 3), (2, 4), (3, 4), (1, 5)]
    demand = {1: {1: 3}, 2: {1: 5}, 4: {1: 2}}
    perm = {1: 50, 2: 10, 3: 30, 4: 20, 5: 40}
    renamed = make_instance({perm[i]: d for i, d in durations.items()},
                            [(perm[a], perm[b]) for a, b in arcs],
                            demand={perm[i]: r for i, r in demand.items()})
    assert flat_indicators(make_instance(durations, arcs, demand)) == flat_indicators(renamed)


def test_case1_values_in_unit_interval():
    import json
    with open(ROOT / "data" / "processed" / "Case1_PSPLIB_j301_1.json", encoding="utf-8") as f:
        ind = flat_indicators(json.load(f))
    assert ind["n"] == 30
    for k in ("os", "sp", "ad", "la", "i5", "tf", "rf", "rs", "rc"):
        assert 0 <= ind[k] <= 1, k
    assert 1 <= ind["ru"] <= 4


@pytest.mark.parametrize("case", ["Case3_nonrenewable.json", "Case5_mp_j90_a2_nr1.json"])
def test_out_of_scope_cases_rejected(case):
    import json
    with open(ROOT / "data" / "processed" / case, encoding="utf-8") as f:
        inst = json.load(f)
    with pytest.raises(IndicatorError):
        network_from_instance(inst)


# Диапазоны guo2021prselection из литобзора

GUO2021_RANGES = {
    "j30": {"cnc": (1.40, 2.07), "os": (0.25, 0.64), "sp": (0.17, 0.41), "ad": (0.11, 0.52),
            "la": (0.00, 0.25), "i5": (0.55, 0.88), "tf": (0.02, 0.50), "ru": (1, 4),
            "rs": (0.14, 1.00), "rc": (0.10, 0.61)},
    "j60": {"cnc": (1.45, 2.08), "os": (0.20, 0.55), "sp": (0.12, 0.27), "ad": (0.16, 0.51),
            "la": (0.00, 0.08), "i5": (0.62, 0.86), "tf": (0.07, 0.50), "ru": (1, 4),
            "rs": (0.14, 1.00), "rc": (0.07, 0.49)},
    "j120": {"cnc": (1.48, 2.09), "os": (0.15, 0.43), "sp": (0.09, 0.18), "ad": (0.19, 0.53),
             "la": (0.01, 0.03), "i5": (0.69, 0.85), "tf": (0.09, 0.44), "ru": (1, 4),
             "rs": (0.09, 0.52), "rc": (0.06, 0.50)},
}
GUO2021_RF = (0.25, 0.50, 0.75, 1.00)
INSTANCES_PER_SET = {"j30": 480, "j60": 480, "j120": 600}

# Таблица дает два знака; расхождение в одну единицу последнего знака допускаем
# (OS J60 max, AD и LA J120 min — см. ADR 0002).
ROUNDING_TOL = 0.01 + 1e-9


@pytest.fixture(scope="module")
def psplib_indicators():
    if not PSPLIB_DIR.exists():
        pytest.skip("PSPLIB не скачан (data/psplib/SOURCES.md)")
    from psplib import iter_set
    return {s: [flat_indicators(inst) for _, inst in iter_set(PSPLIB_DIR / s)]
            for s in GUO2021_RANGES}


@pytest.mark.parametrize("set_name", list(GUO2021_RANGES))
def test_psplib_ranges_match_guo2021(psplib_indicators, set_name):
    rows = psplib_indicators[set_name]
    assert len(rows) == INSTANCES_PER_SET[set_name]
    errors = []
    for key, bounds in GUO2021_RANGES[set_name].items():
        values = [r[key] for r in rows]
        for side, got in ((0, min(values)), (1, max(values))):
            if abs(round(got, 2) - bounds[side]) > ROUNDING_TOL:
                errors.append(f"{key} {'min' if side == 0 else 'max'}: {got:.4f} vs {bounds[side]}")
    assert not errors, errors


@pytest.mark.parametrize("set_name", list(GUO2021_RANGES))
def test_psplib_rf_levels(psplib_indicators, set_name):
    # В таблице — 4 уровня генератора ProGen; посчитанный RF отклоняется от уровня меньше чем на 0.01
    for r in psplib_indicators[set_name]:
        assert min(abs(r["rf"] - lvl) for lvl in GUO2021_RF) < 0.01, r["rf"]


# Построчная сверка с RCPLIB

RCPLIB_FILE = ROOT / "data" / "rcplib" / "RCPLIB (Parameters and BKS).xlsx"
RCPLIB_SUBSETS = {"j30": "J30", "j60": "J60", "j120": "J120"}
RCPLIB_COLUMNS = {"cnc": "CNC", "os": "OS", "sp": "SP", "ad": "AD", "la": "LA", "i5": "I5",
                  "tf": "TF", "rf": "RF", "ru": "RU", "rs": "RS", "rc": "RC"}
# RCPLIB дает три знака; у части значений третий знак усечен, а не округлен (OS j3032_8:
# 0.42299 -> 0.422), поэтому допуск — одна единица третьего знака.
RCPLIB_TOL = 0.001 + 1e-9
# В RCPLIB у j3017_4 три ресурса: R3, который использует только работа 29 (спрос 9 при
# мощности 9), удален. Сетевые индикаторы совпадают, ресурсные — после удаления R3.
RCPLIB_WITHOUT_RESOURCE = {"j3017_4": 3}


@pytest.fixture(scope="module")
def rcplib_rows():
    if not RCPLIB_FILE.exists():
        pytest.skip("RCPLIB не скачан (data/rcplib/SOURCES.md)")
    if not PSPLIB_DIR.exists():
        pytest.skip("PSPLIB не скачан (data/psplib/SOURCES.md)")
    from openpyxl import load_workbook

    wb = load_workbook(RCPLIB_FILE, read_only=True)
    rows = wb["PSPLIB"].iter_rows(values_only=True, max_col=30)
    header = list(next(rows))
    col = {h: header.index(h) for h in ("SubSet", "FileName", *RCPLIB_COLUMNS.values())}
    out = {}
    for r in rows:
        if r[col["FileName"]] is None:
            continue
        name = str(r[col["FileName"]]).removesuffix(".sm")
        out[(r[col["SubSet"]], name)] = {k: float(r[col[c]]) for k, c in RCPLIB_COLUMNS.items()}
    wb.close()
    return out


def drop_resource(instance: dict, resource_id: int) -> dict:
    return {
        "resources": [r for r in instance["resources"] if r["id"] != resource_id],
        "jobs": [{**j, "resource_requirements": [rr for rr in j["resource_requirements"]
                                                 if rr["resource_id"] != resource_id]}
                 for j in instance["jobs"]],
    }


@pytest.mark.parametrize("set_name", list(RCPLIB_SUBSETS))
def test_psplib_matches_rcplib_row_by_row(rcplib_rows, set_name):
    from psplib import iter_set

    checked, errors = 0, []
    for name, inst in iter_set(PSPLIB_DIR / set_name):
        ref = rcplib_rows[(RCPLIB_SUBSETS[set_name], name)]
        if name in RCPLIB_WITHOUT_RESOURCE:
            inst = drop_resource(inst, RCPLIB_WITHOUT_RESOURCE[name])
        got = flat_indicators(inst)
        for key, expected in ref.items():
            if abs(got[key] - expected) > RCPLIB_TOL:
                errors.append(f"{name} {key}: {got[key]:.4f} vs {expected}")
        checked += 1
    assert checked == INSTANCES_PER_SET[set_name]
    assert not errors, errors[:20]


def test_j3017_4_differs_from_rcplib_only_by_resource_r3(rcplib_rows):
    from psplib import parse_sm_file

    inst = parse_sm_file(PSPLIB_DIR / "j30" / "j3017_4.sm")
    ref = rcplib_rows[("J30", "j3017_4")]
    as_is = flat_indicators(inst)
    assert abs(as_is["rf"] - ref["rf"]) > 0.05  # с R3 ресурсные индикаторы расходятся
    for key in ("cnc", "os", "sp", "ad", "la", "i5", "tf"):
        assert abs(as_is[key] - ref[key]) <= RCPLIB_TOL, key
