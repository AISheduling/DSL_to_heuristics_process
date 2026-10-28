from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path

from indicators import compute_indicators
from psplib import instance_name, iter_set, parse_solution_file

ROOT = Path(__file__).resolve().parent.parent
PSPLIB_DIR = ROOT / "data" / "psplib"
RCPLIB_FILE = ROOT / "data" / "rcplib" / "RCPLIB (Parameters and BKS).xlsx"

_RCPLIB_SUBSETS = {"j30": "J30", "j60": "J60", "j90": "J90", "j120": "J120"}


class ReferenceError(ValueError):
    """Эталоны противоречат друг другу (LB > UB, CPM не совпадает и т. п.)."""


@dataclass
class Reference:
    cpm: int                 # Длина критического пути без ресурсов (нижняя граница)
    lb: int                  # Лучшая известная нижняя граница, не меньше cpm
    bks: int                 # Лучшее известное решение (верхняя граница)
    bks_source: str          # "psplib", "rcplib" или "both" — где найдено значение bks
    optimal: bool            # bks доказанно оптимален (lb == bks)

    def as_dict(self) -> dict:
        return asdict(self)


def load_rcplib(path: str | Path = RCPLIB_FILE) -> dict[str, dict[str, int]]:
    """Лист PSPLIB из RCPLIB: {'j301_1': {'cpm', 'lb', 'ub'}} для всех наборов."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        rows = wb["PSPLIB"].iter_rows(values_only=True)
        header = list(next(rows))
        col = {h: header.index(h) for h in ("FileName", "CPM", "LB-lit", "UB-lit")}
        out = {}
        for r in rows:
            if r[col["FileName"]] is None:
                continue
            name = str(r[col["FileName"]]).removesuffix(".sm")
            out[name] = {"cpm": int(r[col["CPM"]]), "lb": int(r[col["LB-lit"]]),
                         "ub": int(r[col["UB-lit"]])}
        return out
    finally:
        wb.close()


def _psplib_solutions(set_name: str, solutions_dir: Path) -> dict[str, dict[str, int]]:
    """{'j301_1': {'opt'?, 'ub', 'lb'?}} из файлов решений PSPLIB одного набора."""
    files = {"opt": f"{set_name}_opt.txt", "ub": f"{set_name}_hrs.txt", "lb": f"{set_name}_lb.txt"}
    out: dict[str, dict[str, int]] = {}
    for kind, fname in files.items():
        path = solutions_dir / fname
        if not path.exists():
            continue
        for (param, inst), value in parse_solution_file(path).items():
            out.setdefault(instance_name(set_name, param, inst), {})[kind] = value
    return out


def load_references(set_name: str, *, psplib_dir: str | Path = PSPLIB_DIR,
                    rcplib: dict[str, dict[str, int]] | None = None,
                    use_rcplib: bool = True) -> dict[str, Reference]:
    """
    Эталоны для всех инстансов набора, ключ — имя инстанса ('j6010_5').
    rcplib: заранее прочитанный load_rcplib() (чтобы не читать Excel на каждый набор);
    если не передан и use_rcplib, читается RCPLIB_FILE, когда он есть.
    """
    psplib_dir = Path(psplib_dir)
    if rcplib is None and use_rcplib and RCPLIB_FILE.exists():
        rcplib = load_rcplib()
    rcplib = rcplib or {}
    solutions = _psplib_solutions(set_name, psplib_dir / "solutions")

    out: dict[str, Reference] = {}
    for name, inst in iter_set(psplib_dir / "instances" / set_name):
        cpm = compute_indicators(inst)["extended"]["cpm_length"]
        sol = solutions.get(name, {})
        rc = rcplib.get(name)
        if rc is not None and rc["cpm"] != cpm:
            raise ReferenceError(f"{name}: CPM {cpm} != RCPLIB {rc['cpm']}")

        uppers = {"psplib": sol.get("opt", sol.get("ub")), "rcplib": rc["ub"] if rc else None}
        uppers = {src: v for src, v in uppers.items() if v is not None}
        if not uppers:
            raise ReferenceError(f"{name}: нет верхней границы ни в PSPLIB, ни в RCPLIB")
        bks = min(uppers.values())
        sources = [src for src, v in uppers.items() if v == bks]
        lb = max([cpm] + [v for v in (sol.get("opt"), sol.get("lb"), rc["lb"] if rc else None)
                          if v is not None])
        if lb > bks:
            raise ReferenceError(f"{name}: нижняя граница {lb} больше BKS {bks}")
        out[name] = Reference(cpm=cpm, lb=lb, bks=bks,
                              bks_source="both" if len(sources) == 2 else sources[0],
                              optimal=lb == bks)
    return out
